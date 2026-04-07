import os
import fitz  # PyMuPDF
import docx
import openpyxl
from logger import log


def extract_content(filepath: str) -> str:
    """Dosya uzantisina gore icerik okur. Hata durumunda bos string doner."""
    ext = os.path.splitext(filepath)[1].lower()
    try:
        if ext == ".pdf":
            return _parse_pdf(filepath)
        elif ext == ".docx":
            return _parse_docx(filepath)
        elif ext == ".xlsx":
            return _parse_xlsx(filepath)
        else:
            return _parse_text(filepath)
    except Exception as e:
        log.error(f"Icerik okuma hatasi ({filepath}): {e}")
        return ""


def _parse_text(filepath: str) -> str:
    """TXT ve kod dosyalarini okur. Farkli encoding'leri dener."""
    encodings = ["utf-8", "cp1254", "latin-1", "iso-8859-9"]
    for enc in encodings:
        try:
            with open(filepath, "r", encoding=enc) as f:
                return f.read()
        except UnicodeDecodeError:
            continue
        except Exception as e:
            log.error(f"TXT okuma hatasi ({filepath}): {e}")
            return ""
    return ""


def _parse_pdf(filepath: str) -> str:
    """PDF dosyalarindan metin cikarir."""
    pages = []
    with fitz.open(filepath) as doc:
        for page in doc:
            pages.append(page.get_text())
    return "\n".join(pages)


def _parse_docx(filepath: str) -> str:
    """Word (DOCX) dosyalarindan metin cikarir."""
    doc = docx.Document(filepath)
    return "\n".join(para.text for para in doc.paragraphs)


def _parse_xlsx(filepath: str) -> str:
    """Excel (XLSX) dosyalarindan metin cikarir."""
    text = []
    wb = openpyxl.load_workbook(filepath, data_only=True, read_only=True)
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            row_text = " ".join(str(cell) for cell in row if cell is not None)
            if row_text.strip():
                text.append(row_text)
    wb.close()
    return "\n".join(text)
